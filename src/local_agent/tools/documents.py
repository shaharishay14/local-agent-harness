"""
Documents as text for files_read: PDF (pypdf), Word .docx, spreadsheets (.xlsx with openpyxl, .csv).

Everything here works on bytes files.py already read through its checked open, so the path checks stay in one
place. A document can come from anyone, so every read is capped (text, pages, sheets, rows, columns, the unzipped
size of a .docx/.xlsx) and files.py runs extract() in a thread with a timeout.

.docx is parsed with the standard library (it's a zip of XML) rather than python-docx, which needs lxml: the
dependencies stay pure Python.
"""

import csv
import io
import re
import zipfile
from xml.etree import ElementTree

FORMATS = {'.pdf': 'pdf', '.docx': 'docx', '.xlsx': 'xlsx', '.csv': 'csv'}
MAX_DOCUMENT_BYTES = 50_000_000
MAX_UNZIPPED_BYTES = 100_000_000  # a .docx/.xlsx that unzips to more is refused: a zip bomb, or not worth it
MAX_TEXT_CHARS = 2_000_000
MAX_PAGES = 500
MAX_SHEETS = 50
MAX_ROWS = 20_000  # per sheet, and for a .csv
MAX_COLUMNS = 200
TIMEOUT_SECONDS = 30


class DocumentError(Exception):
    """A document that can't be read. problem and next_step go to tool_error as they are."""

    def __init__(self, problem: str, next_step: str):
        super().__init__(problem)
        self.problem, self.next_step = problem, next_step


class _Full(Exception):
    """The text reached MAX_TEXT_CHARS."""


class _Text:
    """Lines of extracted text, up to MAX_TEXT_CHARS. limits collects what was left out, for the model."""

    def __init__(self):
        self.lines, self.chars, self.limits = [], 0, []

    def add(self, line: str) -> None:
        if self.chars + len(line) + 1 > MAX_TEXT_CHARS:
            self.lines.append(line[:max(0, MAX_TEXT_CHARS - self.chars)])
            self.limits.append(f'the text stops at {MAX_TEXT_CHARS:,} characters')
            raise _Full
        self.lines.append(line)
        self.chars += len(line) + 1

    def text(self) -> str:
        return '\n'.join(self.lines)


def extract(data: bytes, form: str) -> tuple[str, list[str]]:
    """The text of a document of format form (a FORMATS value), and the limits it hit. Raises DocumentError."""
    out = _Text()
    try:
        {'pdf': _pdf, 'docx': _docx, 'xlsx': _xlsx, 'csv': _csv}[form](data, out)
    except _Full:
        pass
    except DocumentError:
        raise
    except Exception:  # the parsers raise many kinds of errors on a damaged file
        raise DocumentError(f'The file can\'t be read as {form.upper()}: it looks damaged or isn\'t really one.',
                            'Tell the user; they can re-save or re-export it and try again.') from None
    return out.text(), out.limits


def _pdf(data: bytes, out: _Text) -> None:
    from pypdf import PdfReader  # imported here: only a PDF read pays for it

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        # many PDFs are encrypted with an empty password just to set permissions; those open
        try:
            opened = reader.decrypt('')
        except Exception:
            opened = 0
        if not opened:
            raise DocumentError('The PDF is password-protected.',
                                'Ask the user to save a copy without the password and put it in the folder.')
    pages = reader.pages
    if len(pages) > MAX_PAGES:
        out.limits.append(f'only the first {MAX_PAGES} of {len(pages)} pages were read')
    found = False
    for number in range(min(len(pages), MAX_PAGES)):
        text = (pages[number].extract_text() or '').strip()
        found = found or bool(text)
        out.add(f'--- page {number + 1} ---')
        if text:
            out.add(text)
    if not found:
        raise DocumentError('The PDF has no text layer (it is probably a scan or images of pages).',
                            'Tell the user; they can run it through OCR (text recognition) and save the result.')


def _unzipped(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED_BYTES:
        raise DocumentError('The document is too large once unpacked to read safely.',
                            'Tell the user; they can split it or export the part they need.')
    return archive


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
HEADING = re.compile(r'heading\s*(\d)', re.IGNORECASE)


def _paragraph(p) -> str:
    parts = []
    for node in p.iter():
        if node.tag == W + 't':
            parts.append(node.text or '')
        elif node.tag == W + 'tab':
            parts.append('\t')
        elif node.tag in (W + 'br', W + 'cr'):
            parts.append('\n')
    text = ''.join(parts)
    style = p.find(f'{W}pPr/{W}pStyle')
    name = style.get(W + 'val', '') if style is not None else ''
    if text.strip() and (level := HEADING.match(name)):
        return '#' * int(level.group(1)) + ' ' + text
    if text.strip() and name.lower() == 'title':
        return '# ' + text
    return text


def _blocks(container, out: _Text) -> None:
    for child in container:
        if child.tag == W + 'p':
            out.add(_paragraph(child))
        elif child.tag == W + 'tbl':
            for row in child.iter(W + 'tr'):
                cells = [' '.join(_paragraph(p) for p in cell.iter(W + 'p')).strip() for cell in row.iter(W + 'tc')]
                out.add('\t'.join(cells))
        elif child.tag == W + 'sdt':  # a content control: its paragraphs are inside
            content = child.find(W + 'sdtContent')
            if content is not None:
                _blocks(content, out)


def _docx(data: bytes, out: _Text) -> None:
    try:
        archive = _unzipped(data)
    except zipfile.BadZipFile:
        raise DocumentError('The file isn\'t a valid Word document (it may be password-protected, or an old .doc '
                            'renamed).', 'Ask the user to save it as a .docx without a password.') from None
    with archive:
        body = ElementTree.fromstring(archive.read('word/document.xml')).find(W + 'body')
    if body is not None:
        _blocks(body, out)


def _cells(values) -> list[str]:
    cells = ['' if value is None else str(value) for value in values]
    while cells and not cells[-1].strip():
        cells.pop()
    return cells


def _rows(rows, out: _Text, where: str) -> None:
    """Rows of cells, tab-separated, without empty trailing rows and columns, up to MAX_ROWS."""
    blank, wide = 0, False
    for count, values in enumerate(rows):
        if count == MAX_ROWS:
            out.limits.append(f'only the first {MAX_ROWS:,} rows of {where} were read')
            break
        cells = _cells(values[:MAX_COLUMNS])
        if len(values) > MAX_COLUMNS and any(str(v).strip() for v in values[MAX_COLUMNS:] if v is not None) \
                and not wide:
            wide = True
            out.limits.append(f'only the first {MAX_COLUMNS} columns of {where} were read')
        if not cells:
            blank += 1
            continue
        for _ in range(blank):
            out.add('')
        blank = 0
        out.add('\t'.join(cells))


def _xlsx(data: bytes, out: _Text) -> None:
    from openpyxl import load_workbook

    try:
        _unzipped(data).close()
    except zipfile.BadZipFile:
        raise DocumentError('The file isn\'t a valid .xlsx spreadsheet (it may be password-protected, or an old '
                            '.xls renamed).', 'Ask the user to save it as a .xlsx without a password.') from None
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = [sheet for sheet in book.worksheets if hasattr(sheet, 'iter_rows')]  # not chart sheets
        if len(sheets) > MAX_SHEETS:
            out.limits.append(f'only the first {MAX_SHEETS} of {len(sheets)} sheets were read')
        for sheet in sheets[:MAX_SHEETS]:
            out.add(f'--- sheet {sheet.title} ---')
            _rows(sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS + 1, values_only=True), out,
                  f'sheet {sheet.title!r}')
    finally:
        book.close()


def _csv(data: bytes, out: _Text) -> None:
    text = data.decode('utf-8-sig', errors='replace')
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    _rows(csv.reader(io.StringIO(text), dialect), out, 'the file')
