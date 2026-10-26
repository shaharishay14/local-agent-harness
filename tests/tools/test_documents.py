"""
files_read on documents: PDF, Word (.docx) and spreadsheets (.xlsx, .csv) come out as text, in parts like a text
file, within caps, and with errors the model can act on for a locked, damaged or scanned file. The fixtures are
built here, in the test's own workspace.
"""

import io
import json
import time

import docx
import openpyxl
import pytest
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from local_agent.tools import documents
from local_agent.tools.files import file_tools
from tests.scripted import ScriptedModel, call, run_once, say


def make_pdf(pages, password=None) -> bytes:
    """A PDF with one line of text per page ('' for a page with no text), optionally password-protected."""
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                                NameObject('/Subtype'): NameObject('/Type1'),
                                                NameObject('/BaseFont'): NameObject('/Helvetica')}))
    for text in pages:
        page = writer.add_blank_page(612, 792)
        page[NameObject('/Resources')] = DictionaryObject(
            {NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode() if text else b'')
        page[NameObject('/Contents')] = writer._add_object(stream)
    if password is not None:
        writer.encrypt(password, algorithm='RC4-128')
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(path):
    document = docx.Document()
    document.add_heading('Quarterly plan', level=1)
    document.add_paragraph('Ship the harness.')
    document.add_heading('Budget', level=2)
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [('Item', 'Cost'), ('Laptop', '1200')], strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    document.add_paragraph('Thanks, Dana')
    document.save(path)


def make_xlsx(path, sheets):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(path)


@pytest.fixture
def files(home):
    return home.files_dir


async def read(root, path, **arguments):
    tool = {tool.name: tool for tool in file_tools(root)}['files_read']
    return json.loads(await tool.on_invoke_tool(None, json.dumps({'path': path, **arguments})))


# --- each format -------------------------------------------------------------------------------------------

async def test_a_pdf_a_docx_and_an_xlsx_are_read_as_text(files):
    (files / 'docs').mkdir()
    (files / 'docs' / 'report.pdf').write_bytes(make_pdf(['Hello from page one', 'And page two']))
    make_docx(files / 'docs' / 'plan.docx')
    make_xlsx(files / 'docs' / 'budget.xlsx', {'2026': [['Item', 'Cost'], ['Laptop', 1200], [None, None]],
                                                'Notes': [['ok']]})

    pdf = await read(files, 'docs/report.pdf')
    assert pdf['format'] == 'pdf' and pdf['path'] == 'docs/report.pdf' and not pdf['truncated']
    assert pdf['content'] == '--- page 1 ---\nHello from page one\n--- page 2 ---\nAnd page two'

    word = await read(files, 'docs/plan.docx')
    assert word['format'] == 'docx'
    assert word['content'] == '# Quarterly plan\nShip the harness.\n## Budget\nItem\tCost\nLaptop\t1200\nThanks, Dana'

    sheet = await read(files, 'docs/budget.xlsx')
    assert sheet['format'] == 'xlsx'
    assert sheet['content'] == '--- sheet 2026 ---\nItem\tCost\nLaptop\t1200\n--- sheet Notes ---\nok'


async def test_csv_and_text_are_read_and_say_their_format(files):
    (files / 'people.csv').write_text('name,city,,\nDana,Lisbon,,\n,,,\nNoa,Haifa,,\n,,,\n')
    (files / 'semi.csv').write_text('a;b\n1;2\n')
    (files / 'note.md').write_text('# hi\n')
    assert (await read(files, 'people.csv'))['content'] == 'name\tcity\nDana\tLisbon\n\nNoa\tHaifa'
    assert (await read(files, 'semi.csv'))['content'] == 'a\tb\n1\t2'
    note = await read(files, 'note.md')
    assert note['format'] == 'text' and note['content'] == '# hi\n'
    (files / 'REPORT.PDF').write_bytes(make_pdf(['upper']))  # the suffix check ignores case
    assert (await read(files, 'REPORT.PDF'))['format'] == 'pdf'


async def test_a_long_document_is_read_in_parts(files):
    make_xlsx(files / 'big.xlsx', {'data': [[f'row {n}', 'x' * 40] for n in range(200)]})
    first = await read(files, 'big.xlsx')
    assert first['truncated'] and first['next_offset'] == 3500 and len(first['content']) == 3500
    parts, result = [first['content']], first
    while result['truncated']:
        result = await read(files, 'big.xlsx', offset=result['next_offset'])
        parts.append(result['content'])
    text = ''.join(parts)
    assert len(text) == first['total_chars'] and text.endswith('row 199\t' + 'x' * 40)


# --- caps --------------------------------------------------------------------------------------------------

async def test_pages_rows_sheets_and_text_are_capped(files, monkeypatch):
    monkeypatch.setattr(documents, 'MAX_PAGES', 2)
    monkeypatch.setattr(documents, 'MAX_ROWS', 3)
    monkeypatch.setattr(documents, 'MAX_SHEETS', 1)
    monkeypatch.setattr(documents, 'MAX_COLUMNS', 2)
    (files / 'long.pdf').write_bytes(make_pdf(['one', 'two', 'three']))
    pdf = await read(files, 'long.pdf')
    assert 'three' not in pdf['content'] and 'first 2 of 3 pages' in pdf['limits']
    make_xlsx(files / 'wide.xlsx', {'a': [[1, 2, 3]] * 5, 'b': [['hidden']]})
    sheet = await read(files, 'wide.xlsx')
    assert sheet['content'] == '--- sheet a ---\n1\t2\n1\t2\n1\t2'
    assert all(part in sheet['limits'] for part in ['first 1 of 2 sheets', 'first 3 rows', 'first 2 columns'])
    (files / 'rows.csv').write_text('a\n' * 10)
    assert 'first 3 rows' in (await read(files, 'rows.csv'))['limits']

    monkeypatch.setattr(documents, 'MAX_TEXT_CHARS', 50)
    (files / 'flood.csv').write_text('word ' * 20 + '\nmore\n')
    flood = await read(files, 'flood.csv')
    assert flood['total_chars'] <= 50 and 'stops at 50 characters' in flood['limits']


async def test_a_file_too_large_or_too_large_unpacked_is_refused(files, monkeypatch):
    monkeypatch.setattr(documents, 'MAX_DOCUMENT_BYTES', 500)
    (files / 'big.pdf').write_bytes(make_pdf(['x']) + b' ' * 500)
    assert 'too large' in (await read(files, 'big.pdf'))['error']
    monkeypatch.setattr(documents, 'MAX_DOCUMENT_BYTES', 50_000_000)
    monkeypatch.setattr(documents, 'MAX_UNZIPPED_BYTES', 1000)  # what a zip bomb looks like
    make_docx(files / 'plan.docx')
    make_xlsx(files / 'sheet.xlsx', {'a': [[1]]})
    for path in ('plan.docx', 'sheet.xlsx'):
        assert 'too large once unpacked' in (await read(files, path))['error']


async def test_extraction_that_takes_too_long_is_given_up(files, monkeypatch):
    monkeypatch.setattr(documents, 'TIMEOUT_SECONDS', 0.05)
    monkeypatch.setattr(documents, 'extract', lambda data, form: time.sleep(0.5))
    (files / 'slow.pdf').write_bytes(make_pdf(['x']))
    assert 'took too long' in (await read(files, 'slow.pdf'))['error']


# --- documents that can't be read --------------------------------------------------------------------------

async def test_locked_damaged_and_scanned_documents_say_what_to_do(files):
    (files / 'locked.pdf').write_bytes(make_pdf(['secret'], password='pw'))
    (files / 'open.pdf').write_bytes(make_pdf(['readable'], password=''))  # encrypted only for permissions
    (files / 'scan.pdf').write_bytes(make_pdf(['', '']))
    (files / 'broken.pdf').write_bytes(b'%PDF-1.7\n\x00garbage')
    (files / 'broken.docx').write_bytes(b'PK\x03\x04 not really a zip')
    (files / 'old.xlsx').write_bytes(b'\xd0\xcf\x11\xe0 an old binary .xls renamed')

    locked = (await read(files, 'locked.pdf'))['error']
    assert 'password-protected' in locked and 'without the password' in locked and 'secret' not in locked
    assert (await read(files, 'open.pdf'))['content'] == '--- page 1 ---\nreadable'
    assert 'no text layer' in (scan := (await read(files, 'scan.pdf'))['error']) and 'OCR' in scan
    assert 'damaged' in (await read(files, 'broken.pdf'))['error']
    assert 'valid Word document' in (await read(files, 'broken.docx'))['error']
    assert 'valid .xlsx' in (await read(files, 'old.xlsx'))['error']


async def test_other_binary_files_are_still_refused(files):
    (files / 'photo.jpg').write_bytes(b'\xff\xd8\x00\x10JFIF')
    (files / 'data.csv').write_bytes(b'a,b\x00\x00')
    for path in ('photo.jpg', 'data.csv'):
        assert 'binary file' in (await read(files, path))['error']


# --- the path checks still apply ---------------------------------------------------------------------------

async def test_a_document_outside_the_folder_is_not_read(files, tmp_path):
    outside = tmp_path / 'outside'
    outside.mkdir()
    (outside / 'secret.pdf').write_bytes(make_pdf(['the owner secret']))
    (files / 'link.pdf').symlink_to(outside / 'secret.pdf')
    for path in ('../outside/secret.pdf', 'link.pdf', str(outside / 'secret.pdf')):
        result = await read(files, path)
        assert 'error' in result and 'owner secret' not in json.dumps(result)


async def test_a_document_read_in_a_run_is_wrapped_as_third_party(make_runtime, files):
    make_docx(files / 'plan.docx')
    model = ScriptedModel([[call('1', 'files_read', {'path': 'plan.docx'})], [say('It is a quarterly plan.')]])
    await run_once(make_runtime(model), 'what is in plan.docx?')
    output = model.tool_outputs(1)['call_1']
    assert output.startswith('<untrusted_tool_output tool="files_read">') and 'Quarterly plan' in output
